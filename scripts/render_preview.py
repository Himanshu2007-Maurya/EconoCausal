"""Approximate PNG renderer for the deck (no LibreOffice available).

Draws shapes + wrapped text with DejaVu fonts so layout overflow can be
caught visually and programmatically.
"""
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

FD = Path("/usr/share/fonts/truetype/dejavu")
FONTS = {
    (False, False): FD / "DejaVuSans.ttf",
    (True, False): FD / "DejaVuSans-Bold.ttf",
    (False, True): FD / "DejaVuSans.ttf",
    (True, True): FD / "DejaVuSans-Bold.ttf",
}
MONO = FD / "DejaVuSansMono.ttf"
MONO_B = FD / "DejaVuSansMono-Bold.ttf"

SCALE = 1.9  # px per pt-ish


def px(emu):
    return emu / 914400 * 96 * SCALE


_cache = {}


def font(size_pt, bold=False, italic=False, mono=False, name=""):
    s = int(round(size_pt * SCALE * 1.0))
    key = (s, bold, italic, mono)
    if key in _cache:
        return _cache[key]
    if mono:
        p = MONO_B if bold else MONO
    else:
        p = FONTS[(bold, italic)]
    f = ImageFont.truetype(str(p), s)
    _cache[key] = f
    return f


def wrap(draw, text, f, maxw):
    out = []
    for para in text.split("\n"):
        if not para:
            out.append("")
            continue
        words = para.split(" ")
        line = ""
        for w in words:
            trial = (line + " " + w) if line else w
            if draw.textlength(trial, font=f) <= maxw or not line:
                line = trial
            else:
                out.append(line)
                line = w
        out.append(line)
    return out


def render(path, outdir):
    prs = Presentation(path)
    W = int(px(prs.slide_width))
    H = int(px(prs.slide_height))
    outdir.mkdir(parents=True, exist_ok=True)
    warnings = []

    for idx, slide in enumerate(prs.slides, start=1):
        bg = (255, 255, 255)
        try:
            if slide.background.fill.type is not None:
                bg = tuple(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        img = Image.new("RGB", (W, H), bg)
        d = ImageDraw.Draw(img)

        for sh in slide.shapes:
            try:
                l, t = px(sh.left), px(sh.top)
                w, h = px(sh.width), px(sh.height)
            except TypeError:
                continue
            fill = None
            try:
                if sh.fill.type == 1:
                    fill = tuple(sh.fill.fore_color.rgb)
            except Exception:
                pass
            stroke = None
            try:
                if sh.line.color and sh.line.color.rgb is not None:
                    stroke = tuple(sh.line.color.rgb)
            except Exception:
                stroke = None
            st = getattr(sh, "shape_type", None)
            name = str(st) if st is not None else ""
            if fill or stroke:
                if "OVAL" in name:
                    d.ellipse([l, t, l + w, t + h], fill=fill, outline=stroke)
                elif "ROUNDED" in name:
                    r = min(h / 2, h * 0.28)
                    d.rounded_rectangle([l, t, l + w, t + h], radius=r,
                                        fill=fill, outline=stroke)
                else:
                    d.rectangle([l, t, l + w, t + h], fill=fill, outline=stroke)

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            # vertical anchor
            blocks = []
            total_h = 0
            for p in tf.paragraphs:
                if not p.text:
                    continue
                runs = list(p.runs)
                if not runs:
                    continue
                sz = max((r.font.size.pt if r.font.size else 12) for r in runs)
                bold = any(r.font.bold for r in runs)
                italic = any(r.font.italic for r in runs)
                mono = any((r.font.name or "") in ("Consolas", "Courier New")
                           for r in runs)
                color = None
                for r in runs:
                    try:
                        color = tuple(r.font.color.rgb)
                    except Exception:
                        pass
                color = color or (30, 30, 30)
                f = font(sz, bold, italic, mono, name=(r.font.name or ""))
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                lines = wrap(d, p.text, f, max(w, 1) - 1)
                asc = f.getbbox("Ay")[3] - f.getbbox("Ay")[1]
                lh = (sz * SCALE * 1.30) * ls
                blocks.append((lines, f, color, lh, p.alignment))
                total_h += lh * len(lines)

            if tf.vertical_anchor is not None and int(tf.vertical_anchor) == 3:
                cy = t + (h - total_h) / 2
            elif tf.vertical_anchor is not None and int(tf.vertical_anchor) == 2:
                cy = t + h - total_h
            else:
                cy = t
            cy += px(tf.margin_top)
            for lines, f, color, lh, al in blocks:
                for line in lines:
                    lw = d.textlength(line, font=f)
                    if al is not None and int(al) == 2:
                        cx = l + (w - lw) / 2
                    elif al is not None and int(al) == 3:
                        cx = l + w - lw - px(tf.margin_right)
                    else:
                        cx = l + px(tf.margin_left)
                    d.text((cx, cy), line, font=f, fill=color)
                    cy += lh
            if cy > t + h + 2:
                warnings.append(
                    f"slide {idx}: text block overflows by "
                    f"{(cy - (t + h)) / (96 * SCALE):.2f}in :: "
                    f"{tf.text[:60]!r}")
            if l + w > W + 2 or t + h > H + 2:
                warnings.append(f"slide {idx}: shape off-canvas :: "
                                f"{tf.text[:40]!r}")

        img.save(outdir / f"slide-{idx:02d}.png")

    return warnings


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else "EconoCausal_Project_Status_Review.pptx"
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "preview")
    ws = render(p, out)
    print("\n".join(ws) if ws else "no overflow detected")
