"""Shared visual language for the EconoCausal decks.

Palette + small helpers so every slide is built from the same parts.
"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- palette
INK = RGBColor(0x0B, 0x12, 0x20)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
FAINT = RGBColor(0x94, 0xA3, 0xB8)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_L = RGBColor(0x5E, 0xEA, 0xD4)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_L = RGBColor(0xFB, 0xBF, 0x24)
ROSE = RGBColor(0xBE, 0x12, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF8, 0xFA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
TEAL_TINT = RGBColor(0xEC, 0xFD, 0xF5)
AMBER_TINT = RGBColor(0xFF, 0xFB, 0xEB)
ROSE_TINT = RGBColor(0xFF, 0xF1, 0xF2)
GREY_TINT = RGBColor(0xF1, 0xF5, 0xF9)

HEAD = "Segoe UI"
BODY = "Segoe UI"
MONO = "Consolas"

SW, SH = 13.333, 7.5
M = 0.78
CW = SW - 2 * M


class Deck:
    def __init__(self, prs):
        self.prs = prs
        self.page = 0

    # -------------------------------------------------- primitives
    def slide(self, bg=PAPER):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        self.page += 1
        return s

    def tb(self, s, l, t, w, h, text, size=13, bold=False, color=SLATE,
           font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0,
           italic=False, caps=False, tracking=None):
        box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = text.upper() if caps else text
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        if tracking is not None:
            r.font._rPr.set("spc", str(int(tracking * 100)))
        return box

    def rich(self, s, l, t, w, h, chunks, size=13, color=SLATE,
             align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
        box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = spacing
        for text, bold, col, sz in chunks:
            r = p.add_run()
            r.text = text
            r.font.name = BODY
            r.font.size = Pt(sz or size)
            r.font.bold = bold
            r.font.color.rgb = col or color
        return box

    def card(self, s, l, t, w, h, fill=CARD, line=BORDER, radius=0.05):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),
                                Inches(t), Inches(w), Inches(h))
        sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(0.9)
        sh.shadow.inherit = False
        sh.text_frame.word_wrap = True
        return sh

    def rect(self, s, l, t, w, h, fill, line=None, radius=None):
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w),
                                Inches(h))
        if radius:
            sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(0.9)
        sh.shadow.inherit = False
        return sh

    def dot(self, s, l, t, d, color):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t),
                                Inches(d), Inches(d))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def pill(self, s, l, t, text, fg, bg, w=None, size=8.5):
        w = w or (0.18 + 0.078 * len(text))
        sh = self.rect(s, l, t, w, 0.26, bg, None, radius=0.5)
        tf = sh.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = BODY
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = fg
        return sh

    # -------------------------------------------------- composites
    def header(self, s, kicker, title, sub=None):
        self.rect(s, M, 0.445, 0.055, 0.155, TEAL)
        self.tb(s, M + 0.22, 0.44, CW, 0.22, kicker, size=10, bold=True,
                color=TEAL, caps=True, tracking=1.4)
        self.tb(s, M, 0.70, CW, 0.52, title, size=27, bold=True, color=INK,
                font=HEAD, spacing=0.95)
        y = 1.34
        if sub:
            self.tb(s, M, y, CW - 1.2, 0.30, sub, size=12, color=MUTED,
                    spacing=1.15)
            y += 0.34
        self.rect(s, M, y + 0.10, CW, 0.012, BORDER)
        return y + 0.32

    def footer(self, s, dark=False, label="EconoCausal  ·  project brief  ·  "
             "main branch"):
        c = MUTED if not dark else RGBColor(0x7C, 0x8C, 0xA6)
        self.tb(s, M, SH - 0.48, 8.0, 0.24, label, size=8.5, color=c)
        self.tb(s, SW - M - 1.2, SH - 0.48, 1.2, 0.24, f"{self.page:02d}",
                size=8.5, color=c, align=PP_ALIGN.RIGHT, bold=True)

    def label(self, s, l, t, w, text, color=INK, size=9.5):
        self.tb(s, l, t, w, 0.22, text, size=size, bold=True, color=color,
                caps=True, tracking=1.1)
